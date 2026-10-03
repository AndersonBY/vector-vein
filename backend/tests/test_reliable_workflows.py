from __future__ import annotations

import sys
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import openpyxl
import pytest
from docx import Document
from PIL import Image
from vv_llm.types import BackendType

from utilities.file_processing.files import read_file_content
from utilities.file_processing.structured import read_workbook
from utilities.workflow import Workflow
from utilities.workflow.structured_output import make_validator, validate_output
from worker.tasks import output, tools, file_processing
from worker.tasks.llms import base_llm
from worker.tasks.llms.base_llm import BaseLLMTask
from worker.tasks.llms.types.output import ModelOutput


def workflow_data(task_name, **values):
    return {"nodes": [{"id": "node", "type": "Test", "category": "tools", "data": {
        "task_name": task_name, "template": {key: {"value": value} for key, value in values.items()}
    }}], "edges": []}


def field(data, name):
    return Workflow(data).get_node_field_value("node", name)


@pytest.fixture
def isolated_nodes(monkeypatch, tmp_path):
    statuses = []
    monkeypatch.setattr(Workflow, "report_node_status", lambda *args: True)
    monkeypatch.setattr(Workflow, "report_workflow_status", lambda self, status, *args: statuses.append(status))
    monkeypatch.setattr(output, "Settings", lambda: SimpleNamespace(output_folder=str(tmp_path)))
    monkeypatch.setattr(tools, "Settings", lambda: SimpleNamespace(output_folder=str(tmp_path)))
    return statuses


class MemoryCache(dict):
    def set(self, key, value, expire=None):
        self[key] = value


def model_task(monkeypatch, responses, prompts=None, schema=None, store=None):
    task: Any = BaseLLMTask.__new__(BaseLLMTask)
    task.node_id = "node"
    task.workflow = Workflow(workflow_data("llms.local_llm"))
    task.MODEL_TYPE = BackendType.Local
    task.model = "test-model"
    task.model_settings = SimpleNamespace(endpoints=[], model_dump=lambda **kwargs: {"id": "test-model"})
    task.prompts = prompts or ["extract record"]
    task.input_prompt = task.prompts
    task.prompts_count = len(task.prompts)
    task.system_prompt = "Use only the provided source."
    task.output_validator = make_validator(schema)
    task.max_repairs = 1
    task.cache_results = store is not None
    task.cache_version = ""
    for name in ("temperature", "top_p", "thinking", "reasoning_effort", "response_format", "tools", "tool_choice"):
        setattr(task, name, None)
    task.extra_body = {}
    task.stream = False
    task.use_function_call = False
    task.content_outputs = [""] * len(task.prompts)
    task.reasoning_content_outputs = [""] * len(task.prompts)
    task.function_call_outputs = [[] for _ in task.prompts]
    task.function_call_arguments_batches = [{} for _ in task.prompts]
    task.total_prompt_tokens = task.total_completion_tokens = 0
    task.get_max_concurrent_requests = lambda: 1
    calls = []
    def process(prompt, index):
        calls.append((prompt, index))
        item = responses.pop(0)
        if isinstance(item, Exception):
            raise item
        return ModelOutput(content_output=item, prompt_tokens=5, completion_tokens=3)
    task.process_prompt = process
    if store is not None:
        monkeypatch.setattr(base_llm, "cache", store)
    return task, calls


SCHEMA = {"type": "object", "properties": {"count": {"type": "integer", "minimum": 0}}, "required": ["count"], "additionalProperties": False}


def test_failed_model_item_blocks_batch_and_rerun_reuses_only_successes(monkeypatch):
    store = MemoryCache()
    task, _ = model_task(monkeypatch, ['{"count": 0}', RuntimeError("offline")], ["first", "second"], SCHEMA, store)
    with pytest.raises(RuntimeError, match="indexes: 1"):
        task.run()
    assert len(store) == 1
    assert field(task.workflow.data, "run_stats")["failures"] == [{"index": 1, "error": "RuntimeError"}]
    rerun, calls = model_task(monkeypatch, ['{"count": 2}'], ["first", "second"], SCHEMA, store)
    result = rerun.run()
    assert [index for _, index in calls] == [1]
    assert field(result, "output") == ['{"count": 0}', '{"count": 2}']
    assert field(result, "run_stats")["cache_hits"] == 1
    changed, calls = model_task(monkeypatch, ['{"count": 3}'], ["first"], SCHEMA, store)
    changed.cache_version = "new"
    changed.run()
    assert len(calls) == 1


def test_schema_repairs_are_bounded_and_count_usage(monkeypatch):
    task, calls = model_task(monkeypatch, ['{"count": "wrong"}', '{"count": 2}'], schema=SCHEMA)
    data = task.run()
    assert len(calls) == 2 and "Validation error" in calls[1][0]
    assert field(data, "run_stats")["prompt_tokens"] == 10
    failing, calls = model_task(monkeypatch, ["invalid", "still invalid"], schema=SCHEMA)
    with pytest.raises(RuntimeError):
        failing.run()
    assert len(calls) == 2


def test_schema_references_are_local_and_nonfinite_json_is_rejected():
    validator = make_validator({"$defs": {"count": {"type": "integer"}}, "$ref": "#/$defs/count"})
    assert validate_output("2", validator) == 2
    with pytest.raises(ValueError):
        validate_output("NaN", make_validator({}))
    with pytest.raises(Exception):
        validate_output("2", make_validator({"$ref": "https://example.invalid/schema.json"}))


def test_python_failure_marks_workflow_failed_and_restores_stdout(isolated_nodes):
    original_stdout = sys.stdout
    data = workflow_data("tools.programming_function", language="python", list_input=False,
                         code="def main():\n    raise ValueError('sample failure')")
    with pytest.raises(RuntimeError, match="indexes"):
        tools.programming_function(data, "node")
    assert isolated_nodes == [500]
    assert sys.stdout is original_stdout
    allowed = workflow_data("tools.programming_function", language="python", list_input=False,
                            continue_on_error=True, code="def main():\n    raise ValueError('sample failure')")
    result = tools.programming_function(allowed, "node")
    assert "sample failure" in field(result, "error_msg")


def test_source_workbook_retains_all_sheets_and_coordinates(tmp_path, isolated_nodes):
    file = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    assert wb.active is not None
    wb.active.title = "Orders"
    wb.active.append(["name", "count"])
    wb.active.append(["A,B", 0])
    ws = wb.create_sheet("Checks")
    ws.append(["date", "passed"])
    ws.append([date(2025, 5, 10), False])
    wb.save(file)
    text = read_file_content(file)
    assert "## Orders" in text and "## Checks" in text and '"A,B",0' in text
    result = file_processing.file_loader(workflow_data("file_processing.file_loader", files=[str(file)], output_format="structured"), "node")
    data = field(result, "output")
    assert data["source"] == str(file.resolve())
    assert data["sheets"][1]["rows"][0] == {"row_number": 2, "values": ["2025-05-10T00:00:00", False]}
    assert len(read_workbook(file, ["Checks"])["sheets"]) == 1
    with pytest.raises(ValueError):
        read_workbook(file, ["Missing"])


def export_node(tmp_path, name, extension, content, **kwargs):
    result = output.document(workflow_data("output.document", file_name=name, export_type=extension,
                                        content=content, content_format="structured", **kwargs), "node")
    return Path(field(result, "output"))


def test_neutral_report_exports_consistent_workbook_docx_and_png(tmp_path, isolated_nodes):
    rows = [["2025-05-10", 0], ["2025-05-11", 2]]
    workbook = export_node(tmp_path, "report", ".xlsx", {"sheets": [
        {"name": "Daily", "columns": ["Date", "Count"], "rows": rows, "column_types": ["date", "integer"]},
        {"name": "Checks", "columns": ["Passed", "Note"], "rows": [[False, "=literal"]]}
    ]})
    chart = export_node(tmp_path, "counts", ".png", {"type": "bar", "title": "Daily counts", "x": [row[0] for row in rows], "ylabel": "Records", "series": [{"name": "Count", "y": [row[1] for row in rows]}]})
    report = export_node(tmp_path, "report", ".docx", {"title": "Sample data report", "sections": [{"heading": "Results", "paragraphs": ["Total records: 2"], "table": {"columns": ["Date", "Count"], "rows": rows}, "images": [{"path": str(chart), "caption": "Figure 1: Daily counts"}]}]})
    wb = openpyxl.load_workbook(workbook)
    assert wb.sheetnames == ["Daily", "Checks"]
    assert wb["Daily"]["B2"].value == 0 and wb["Checks"]["A2"].value is False
    assert wb["Checks"]["B2"].data_type == "s"
    wb.close()
    doc = Document(str(report))
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1
    assert doc.tables[0].cell(2, 1).text == "2"
    with Image.open(chart) as image:
        assert image.format == "PNG" and image.width > 100


def test_export_failure_preserves_existing_file_and_removes_temporary(tmp_path, isolated_nodes):
    target = tmp_path / "report.xlsx"
    target.write_bytes(b"previous file")
    with pytest.raises(ValueError, match="Row width"):
        export_node(tmp_path, "report", ".xlsx", {"sheets": [{"name": "Data", "columns": ["count"], "rows": [[1, 2]]}]}, overwrite=True)
    assert target.read_bytes() == b"previous file"
    assert list(tmp_path.iterdir()) == [target]


def test_csv_export_preserves_quoted_commas(tmp_path, isolated_nodes):
    result = output.document(workflow_data("output.document", file_name="csv", export_type=".xlsx", content='name,count\n"A,B",0'), "node")
    wb = openpyxl.load_workbook(field(result, "output"))
    assert wb.active is not None
    assert wb.active["A2"].value == "A,B"
    wb.close()


def test_python_top_level_executes_once_and_batch_lengths_are_checked(isolated_nodes):
    result = tools.programming_function(workflow_data("tools.programming_function", language="python", code="print('once')"), "node")
    assert field(result, "console_msg") == "once\n"
    data = workflow_data("tools.programming_function", language="python", list_input=True,
                         code="def main(left, right): return left + right", left=[1, 2], right=[3])
    with pytest.raises(ValueError, match="equal lengths"):
        tools.programming_function(data, "node")


@pytest.mark.parametrize("exhausted", [False, True])
def test_poll_retry_preserves_state_and_marks_exhaustion_failed(monkeypatch, isolated_nodes, exhausted):
    from celery.exceptions import MaxRetriesExceededError, Retry
    from worker.tasks import task, TaskRetry

    @task(max_retries=2)
    def poll(workflow_data, node_id):
        workflow_data["poll_count"] = 1
        raise TaskRetry("poll", {**workflow_data, "node_id": node_id}, 1)

    calls = []
    def retry(**kwargs):
        calls.append(kwargs)
        raise MaxRetriesExceededError() if exhausted else Retry()
    monkeypatch.setattr(poll.celery_task, "retry", retry)
    with pytest.raises(MaxRetriesExceededError if exhausted else Retry):
        poll(workflow_data("tools.poll"), "node")
    assert calls[0]["args"][0]["poll_count"] == 1
    assert "node_id" not in calls[0]["args"][0]
    assert calls[0]["max_retries"] == 2
    assert isolated_nodes == ([500] if exhausted else [])
