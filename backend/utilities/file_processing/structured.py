"""Structured, source-preserving workbook input and report exports."""
from __future__ import annotations

import csv
import io
import math
from datetime import date, datetime, time
from pathlib import Path
from threading import RLock

import openpyxl
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches


# ponytail: serialize rendering; use isolated workers if chart throughput matters.
_CHART_LOCK = RLock()


def read_workbook(file: str | Path, sheet_names: list[str] | None = None) -> dict:
    workbook = openpyxl.load_workbook(file, read_only=True, data_only=False)
    try:
        names = sheet_names or workbook.sheetnames
        if len(set(names)) != len(names) or any(name not in workbook.sheetnames for name in names):
            raise ValueError("Worksheet names must be unique and exist in the workbook")
        sheets = []
        for name in names:
            rows = workbook[name].iter_rows(values_only=True)
            columns = [value.isoformat() if isinstance(value, (date, datetime, time)) else value for value in next(rows, ())]
            records = []
            for number, row in enumerate(rows, 2):
                values = [value.isoformat() if isinstance(value, (date, datetime, time)) else value for value in row]
                records.append({"row_number": number, "values": values})
            sheets.append({"name": name, "columns": columns, "rows": records})
        return {"source": str(Path(file).resolve()), "sheets": sheets}
    finally:
        workbook.close()


def workbook_text(data: dict) -> str:
    parts = []
    for sheet in data["sheets"]:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(sheet["columns"])
        writer.writerows(row["values"] for row in sheet["rows"])
        heading = f"## {sheet['name']}\n" if len(data["sheets"]) > 1 else ""
        parts.append(heading + output.getvalue())
    return "\n".join(parts)


def _cell_value(value, kind: str | None):
    if value is None or kind in (None, "auto"):
        return value
    if kind == "string":
        return str(value)
    if kind == "integer":
        if isinstance(value, bool) or int(value) != float(value):
            raise ValueError("Expected an integer")
        return int(value)
    if kind == "number":
        if isinstance(value, bool) or not math.isfinite(float(value)):
            raise ValueError("Expected a finite number")
        return float(value)
    if kind == "boolean":
        if not isinstance(value, bool):
            raise ValueError("Expected a JSON boolean")
        return value
    if kind == "date":
        value = date.fromisoformat(value) if isinstance(value, str) else value
        if isinstance(value, datetime):
            value = value.date()
        if not isinstance(value, date):
            raise ValueError("Expected an ISO date")
        return value
    if kind == "datetime":
        value = datetime.fromisoformat(value) if isinstance(value, str) else value
        if not isinstance(value, datetime) or value.tzinfo is not None:
            raise ValueError("Expected an ISO datetime without a timezone")
        return value
    raise ValueError(f"Unsupported column type: {kind}")


def write_workbook(data: dict, target: Path) -> None:
    sheets = data.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("Provide a non-empty sheets array")
    workbook = openpyxl.Workbook()
    initial = workbook.active
    assert initial is not None
    workbook.remove(initial)
    names = set()
    for sheet in sheets:
        name = sheet["name"]
        if not isinstance(name, str) or not name or len(name) > 31 or name.casefold() in names:
            raise ValueError("Worksheet names must be unique and 1-31 characters long")
        names.add(name.casefold())
        columns = sheet["columns"]
        if not isinstance(columns, list) or not columns:
            raise ValueError("Provide a non-empty columns array")
        ws = workbook.create_sheet(name)
        kinds = sheet.get("column_types", [None] * len(columns))
        formats = sheet.get("column_formats", [None] * len(columns))
        widths = sheet.get("column_widths", [20] * len(columns))
        if any(len(values) != len(columns) for values in (kinds, formats, widths)):
            raise ValueError("Column settings must match the number of columns")
        ws.append(columns)
        if not isinstance(sheet["rows"], list):
            raise ValueError("Worksheet rows must be an array")
        for row in sheet["rows"]:
            values = row["values"] if isinstance(row, dict) else row
            if not isinstance(values, (list, tuple)) or len(values) != len(columns):
                raise ValueError(f"Row width does not match columns in worksheet {name}")
            ws.append([_cell_value(value, kind) for value, kind in zip(values, kinds)])
        for column, width in enumerate(widths, 1):
            if not isinstance(width, (int, float)) or not math.isfinite(width) or width <= 0:
                raise ValueError("Column widths must be positive numbers")
            ws.column_dimensions[get_column_letter(column)].width = width
        for row in ws:
            for cell in row:
                # Structured exports contain data; strings must not become formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row > 1 and formats[cell.column - 1]:
                    cell.number_format = formats[cell.column - 1]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    workbook.save(target)
    workbook.close()


def write_report(data: dict, target: Path, template_file: str | None = None) -> None:
    document = Document(template_file) if template_file else Document()
    if data.get("title"):
        document.add_heading(data["title"], 0)
    for section in data.get("sections", []):
        if section.get("heading"):
            document.add_heading(section["heading"], 1)
        for paragraph in section.get("paragraphs", []):
            document.add_paragraph(paragraph)
        if section.get("table"):
            spec = section["table"]
            columns = spec["columns"]
            table = document.add_table(rows=1, cols=len(columns))
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, columns):
                cell.text = str(value)
            for values in spec["rows"]:
                if len(values) != len(columns):
                    raise ValueError("Report table row width does not match columns")
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = "" if value is None else str(value)
        for image in section.get("images", []):
            document.add_picture(str(image["path"]), width=Inches(6))
            if image.get("caption"):
                document.add_paragraph(image["caption"])
    document.save(str(target))


def write_chart(data: dict, target: Path) -> None:
    # Object-oriented Agg rendering works without a browser or GUI event loop.
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    with _CHART_LOCK:
        kind = data.get("type", "line")
        if kind not in ("line", "bar", "scatter"):
            raise ValueError("Chart type must be line, bar or scatter")
        x = data["x"]
        series = data["series"]
        if not x or not series:
            raise ValueError("Chart data must not be empty")
        figure = Figure(figsize=(9, 5), layout="constrained")
        FigureCanvasAgg(figure)
        axes = figure.subplots()
        for index, item in enumerate(series):
            y = item["y"]
            if len(y) != len(x) or any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in y):
                raise ValueError("Chart series must contain one finite number per x value")
            if kind == "bar":
                width = .8 / len(series)
                axes.bar([i + (index - (len(series) - 1) / 2) * width for i in range(len(x))], y, width=width, label=item["name"])
                axes.set_xticks(range(len(x)), [str(value) for value in x])
            else:
                getattr(axes, "plot" if kind == "line" else "scatter")(x, y, label=item["name"])
        axes.set(title=data.get("title", ""), xlabel=data.get("xlabel", ""), ylabel=data.get("ylabel", ""))
        axes.legend()
        figure.savefig(target, format="png", dpi=150)
