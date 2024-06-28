# -*- coding: utf-8 -*-
# @Author: Bi Ying
# @Date:   2023-04-26 21:10:52
# @Last Modified by:   Bi Ying
# @Last Modified time: 2024-06-28 19:37:00
import uuid
from io import StringIO
from pathlib import Path
from datetime import datetime

import yagmail
import openpyxl
import markdown2
import pandas as pd
from docx import Document
from docx.oxml.ns import qn

from worker.tasks import task, timer
from utilities.general import mprint
from utilities.config import Settings
from utilities.workflow import Workflow
from utilities.media_processing import TTSClient
from utilities.file_processing import HtmlToDocx, process_pdf, static_file_server


@task
@timer
def text(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    text: str = workflow.get_node_field_value(node_id, "text")
    workflow.get_node_field_value(node_id, "output_title")
    workflow.get_node_field_value(node_id, "render_markdown")
    workflow.update_node_field_value(node_id, "text", text)
    workflow.update_node_field_value(node_id, "output", text)
    return workflow.data


@task
@timer
def table(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    content: str = workflow.get_node_field_value(node_id, "content")
    content_type: str = workflow.get_node_field_value(node_id, "content_type")
    workflow.get_node_field_value(node_id, "show_table")

    if len(content) == 0:
        result = ""
    elif content_type == "file":
        df = pd.read_csv(content)
        result = df.to_json(orient="records")
    elif content_type == "csv":
        df = pd.read_csv(StringIO(content))
        result = df.to_json(orient="records")
    elif content_type == "json":
        result = content

    workflow.update_node_field_value(node_id, "output", result)
    return workflow.data


@task
@timer
def email(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    to_email: str = workflow.get_node_field_value(node_id, "to_email")
    subject: str = workflow.get_node_field_value(node_id, "subject")
    content_html: str = workflow.get_node_field_value(node_id, "content_html")
    settings = Settings()
    yag = yagmail.SMTP(
        user=settings.get("email.user"),
        password=settings.get("email.password"),
        host=settings.get("email.smtp_host"),
        port=settings.get("email.smtp_port"),
        smtp_ssl=settings.get("email.smtp_ssl"),
    )
    email_send_result = yag.send(to_email, subject, [content_html])
    mprint("email_send_result", email_send_result)
    return workflow.data


@task
@timer
def document(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    output_folder = Path(Settings().output_folder)
    file_name = workflow.get_node_field_value(node_id, "file_name")
    content = workflow.get_node_field_value(node_id, "content")
    workflow.get_node_field_value(node_id, "show_download")
    contents = [content]
    export_type = workflow.get_node_field_value(node_id, "export_type")

    local_file = output_folder / f"{file_name}{export_type}"
    if local_file.exists():
        local_file = output_folder / f"{file_name}_{datetime.now().strftime('%Y%m%d%H%M%S')}{export_type}"
    if export_type.endswith((".txt", ".md", ".html", ".json", ".csv")):
        with open(local_file, "w") as txt_file:
            txt_file.write("\n".join(contents))
    elif export_type.endswith(".docx"):
        content_str = "\n".join(contents)
        html_content = markdown2.markdown(content_str)
        new_parser = HtmlToDocx()
        document = Document()
        new_parser.add_html_to_document(html_content, document)
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.name = "微软雅黑"
                r = run._element.rPr.rFonts
                r.set(qn("w:eastAsia"), "微软雅黑")

        document.save(local_file)
    elif export_type.endswith(".xlsx"):
        content_str = "\n".join(contents)
        lines = content_str.split("\n")
        wb = openpyxl.Workbook()
        ws = wb.active
        for line in lines:
            ws.append(line.split(","))
        wb.save(local_file)

    file_full_path = str(local_file.resolve())
    workflow.update_node_field_value(node_id, "output", file_full_path)
    return workflow.data


@task
@timer
def audio(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    audio_type = workflow.get_node_field_value(node_id, "audio_type", "text_to_speech")
    content = workflow.get_node_field_value(node_id, "content")
    file_link = workflow.get_node_field_value(node_id, "file_link")
    direct_play = workflow.get_node_field_value(node_id, "direct_play")
    output_type = workflow.get_node_field_value(node_id, "output_type")

    output_file_path = ""

    if audio_type == "text_to_speech":
        tts_provider = workflow.get_node_field_value(node_id, "tts_provider")
        tts_model = workflow.get_node_field_value(node_id, "tts_model")
        tts_voice = workflow.get_node_field_value(node_id, "tts_voice")

        tts_client = TTSClient(provider=tts_provider, model=tts_model)
        if direct_play:
            tts_client.stream(text=content, voice=tts_voice)
        else:
            output_folder = Path(Settings().output_folder)
            output_file_path = tts_client.create(text=content, voice=tts_voice, output_folder=output_folder)
            output_file_path = static_file_server.get_file_url(output_file_path)
    elif audio_type == "play_audio":
        output_file_path = file_link
    else:
        mprint(f"不支持的音频类型: {audio_type}")
        output_file_path = file_link

    workflow.update_node_field_value(node_id, "audio_url", output_file_path)
    if output_type == "only_link":
        workflow.update_node_field_value(node_id, "output", output_file_path)
    elif output_type == "markdown":
        workflow.update_node_field_value(node_id, "output", f"[{output_file_path}]({output_file_path})")
    elif output_type == "html":
        workflow.update_node_field_value(node_id, "output", f'<a href="{output_file_path}">{output_file_path}</a>')

    return workflow.data


@task
@timer
def mindmap(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    workflow.get_node_field_value(node_id, "content")
    workflow.get_node_field_value(node_id, "show_mind_map")
    return workflow.data


@task
@timer
def mermaid(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    workflow.get_node_field_value(node_id, "content")
    workflow.get_node_field_value(node_id, "show_mermaid")
    return workflow.data


@task
@timer
def echarts(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    workflow.get_node_field_value(node_id, "option")
    workflow.get_node_field_value(node_id, "show_echarts")
    return workflow.data


@task
@timer
def workflow_invoke_output(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    workflow.get_node_field_value(node_id, "value")
    return workflow.data


@task
@timer
def picture_render(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    render_type = workflow.get_node_field_value(node_id, "render_type")
    input_content = workflow.get_node_field_value(node_id, "content")
    output_type = workflow.get_node_field_value(node_id, "output_type")

    if isinstance(input_content, str):
        contents = [input_content]
    elif isinstance(input_content, list):
        contents = input_content

    static_path = static_file_server.static_folder_path
    image_path = static_path / "images" / "pdf_render" / uuid.uuid4().hex

    results = []
    for content in contents:
        if render_type == "pdf":
            image_files = process_pdf(content, "render_images", image_path)
            image_urls = [
                static_file_server.get_file_url(Path(image_file).relative_to(static_path.absolute()).as_posix())
                for image_file in image_files
            ]
            if output_type == "only_link":
                results.append(image_urls)
            elif output_type == "markdown":
                results.append([f"![{url}]({url})" for url in image_urls])
            elif output_type == "html":
                results.append([f'<img src="{url}"/>' for url in image_urls])

    output = results[0] if isinstance(input_content, str) else results
    workflow.update_node_field_value(node_id, "output", output)
    return workflow.data


@task
@timer
def html(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)
    html_code = workflow.get_node_field_value(node_id, "html_code")
    output_folder = Path(Settings().output_folder)
    file_name = f"{uuid.uuid4().hex}.html"
    file_path = output_folder / file_name
    with open(file_path, "w", encoding="utf8") as f:
        f.write(html_code)
    html_url = static_file_server.get_file_url(file_path)
    workflow.update_node_field_value(node_id, "output", html_url)
    return workflow.data
