"""Run a neutral report DAG against a local deterministic model stub.

No external model or credentials are used. The stub deliberately returns one
invalid result, exercising schema repair; a second run must use the result cache.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import threading
import hashlib
from importlib.metadata import version
from copy import deepcopy
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, cast

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))


def node(identity, kind, category, task_name, values, outputs=("output",)):
    template = {}
    for key, value in values.items():
        template[key] = {"value": value, "type": "dict" if isinstance(value, dict) else "str", "name": key,
                         "display_name": key, "show": False, "field_type": "textarea", "list": False}
    for key in outputs:
        template[key] = {"value": "", "is_output": True, "type": "str", "field_type": "", "show": True, "name": key, "display_name": key}
    return {"id": identity, "type": kind, "category": category, "position": {"x": 0, "y": 0},
            "data": {"task_name": task_name, "has_inputs": True, "template": template}}


def create_workflow(source: Path) -> dict:
    schema = {"type": "object", "properties": {"count": {"type": "integer", "minimum": 0}, "summary": {"type": "string"}},
              "required": ["count", "summary"], "additionalProperties": False}
    report_code = """import json

def main(source, draft):
    result = json.loads(draft)
    rows = [row['values'] for row in source['sheets'][0]['rows']]
    total = sum(row[1] for row in rows)
    if result['count'] != total:
        raise ValueError('Model total differs from source records')
    return {
        'xlsx': source,
        'docx': {'title': 'Sample data report', 'sections': [{'heading': 'Verified results',
            'paragraphs': [result['summary'], f'Total records: {total}'],
            'table': {'columns': ['Date', 'Count'], 'rows': rows}}]},
        'png': {'type': 'bar', 'title': 'Daily counts', 'xlabel': 'Date', 'ylabel': 'Records',
            'x': [row[0] for row in rows], 'series': [{'name': 'Count', 'y': [row[1] for row in rows]}]}
    }
"""
    nodes = [
        node("load", "FileLoader", "fileProcessing", "file_processing.file_loader", {"files": [str(source)], "output_format": "structured"}),
        node("prompt", "ProgrammingFunction", "tools", "tools.programming_function", {"language": "python", "source": {},
             "code": "import json\ndef main(source): return 'Sum the Count column. Use only these records: ' + json.dumps(source, sort_keys=True)"}),
        node("extract", "LocalLLM", "llms", "llms.local_llm", {"prompt": "", "llm_model": "report-demo", "temperature": 0,
             "system_prompt": "Return facts from the supplied workbook.", "output_schema": schema, "max_repairs": 1,
             "cache_results": True, "cache_version": "report-demo-v1", "stream": False}),
        node("prepare", "ProgrammingFunction", "tools", "tools.programming_function", {"language": "python", "source": {}, "draft": "", "code": report_code}),
        node("split", "JsonProcess", "controlFlows", "control_flows.json_process", {"input": {}, "process_mode": "get_multiple_values", "keys": ["xlsx", "docx", "png"]},
             outputs=("output-xlsx", "output-docx", "output-png")),
    ]
    connections = [("load", "output", "prompt", "source"), ("prompt", "output", "extract", "prompt"),
                   ("load", "output", "prepare", "source"), ("extract", "output", "prepare", "draft"),
                   ("prepare", "output", "split", "input")]
    for extension in ("xlsx", "docx", "png"):
        nodes.append(node(extension, "Document", "outputs", "output.document", {"file_name": "neutral-report", "export_type": "." + extension,
                          "content": {}, "content_format": "structured", "overwrite": True}))
        connections.append(("split", "output-" + extension, extension, "content"))
    for index, item in enumerate(nodes):
        item["position"] = {"x": index * 320, "y": 0 if index < 5 else (index - 5) * 220}
    edges = [{"id": str(index), "source": a, "sourceHandle": out, "target": b, "targetHandle": field} for index, (a, out, b, field) in enumerate(connections)]
    return {"nodes": nodes, "edges": edges}


class StubModel(BaseHTTPRequestHandler):
    requests = 0

    def log_message(self, format: str, *args):
        pass

    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        type(self).requests += 1
        content = {"count": "invalid" if type(self).requests == 1 else 2, "summary": "Two sample records."}
        response = {"id": "demo", "object": "chat.completion", "created": 0, "model": "report-demo",
                    "choices": [{"index": 0, "message": {"role": "assistant", "content": json.dumps(content)}, "finish_reason": "stop"}],
                    "usage": {"prompt_tokens": 20, "completion_tokens": 10, "total_tokens": 30}}
        encoded = json.dumps(response).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=BACKEND / ".cache/reliable-report-demo")
    args = parser.parse_args()
    destination = args.output_dir.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    runtime = destination / "runtime"
    runtime.mkdir(exist_ok=True)
    os.chdir(runtime)

    import openpyxl
    from docx import Document
    from PIL import Image
    from models import create_tables, Setting, Workflow as WorkflowModel, WorkflowRunRecord, database
    from utilities.config import DEFAULT_SETTINGS, cache, config
    from utilities.workflow import Workflow
    from celery_worker import app
    from background_task.workflow_tasks import run_workflow

    create_tables()
    server = ThreadingHTTPServer(("127.0.0.1", 0), StubModel)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        settings = cast(dict[str, Any], deepcopy(DEFAULT_SETTINGS))
        settings["output_folder"] = str(destination)
        settings["llm_settings"]["endpoints"].append({"id": "demo", "api_base": f"http://127.0.0.1:{server.server_port}/v1", "api_key": "not-needed", "concurrent_requests": 1})
        settings["llm_settings"]["backends"]["local"]["models"]["report-demo"] = {
            "id": "report-demo", "context_length": 8192, "max_output_tokens": 512,
            "endpoints": ["demo"], "function_call_available": False, "response_format_available": True,
        }
        Setting.create(data=settings)
        app.conf.update(task_always_eager=True, task_eager_propagates=True)
        source = destination / "input.xlsx"
        wb = openpyxl.Workbook()
        assert wb.active is not None
        wb.active.title = "Daily"
        for row in [["Date", "Count"], ["2025-05-10", 0], ["2025-05-11", 2]]:
            wb.active.append(row)
        ws = wb.create_sheet("Checks")
        ws.append(["Source", "Passed"])
        ws.append(["sample", True])
        wb.save(source)
        wb.close()
        payload = create_workflow(source)
        (destination / "workflow.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        model = WorkflowModel.create(title="Neutral report demonstration", data=payload)
        records = []
        for _ in range(2):
            data = deepcopy(payload)
            data["wid"] = model.wid.hex
            record = WorkflowRunRecord.create(workflow=model, data=data, status="RUNNING")
            data["rid"] = record.rid.hex
            result = run_workflow.run(data).get()
            saved = WorkflowRunRecord.get_by_id(record.rid)
            assert saved.status == "FINISHED", saved.status
            records.append({"rid": record.rid.hex, "status": saved.status,
                            "model": Workflow(result).get_node_field_value("extract", "run_stats")})
        assert StubModel.requests == 2, StubModel.requests
        assert records[1]["model"]["cache_hits"] == 1
        wb = openpyxl.load_workbook(destination / "neutral-report.xlsx")
        assert wb.sheetnames == ["Daily", "Checks"] and wb["Daily"]["B3"].value == 2
        wb.close()
        doc = Document(str(destination / "neutral-report.docx"))
        assert doc.tables[0].cell(2, 1).text == "2"
        with Image.open(destination / "neutral-report.png") as image:
            image.verify()
        manifest = {"python": sys.version.split()[0],
                    "packages": {name: version(name) for name in ("vv-llm", "jsonschema", "matplotlib", "openpyxl", "python-docx", "celery")},
                    "input_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                    "model_transport": "local deterministic stub", "model_quality_evaluated": False,
                    "requests": StubModel.requests, "runs": records, "artifacts": ["neutral-report.xlsx", "neutral-report.docx", "neutral-report.png"]}
        (destination / "verification.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        print(json.dumps(manifest, indent=2))
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
        cache.close()
        database.close()
        config.close()


if __name__ == "__main__":
    main()
